import pandas as pd
import glob
import numpy as np

files = glob.glob('../data/*-data.csv.gz')


def excel_kills_accesibility():
    ''' Why do people insist on using proprietary formats '''
    ''' Oh it gets better - outdated proprietary formats: 
    
    InvalidFileException: openpyxl does not support the old .xls file format, please use xlrd to read this file, or convert it to the more recent .xlsx file format.
    
    '''
    from openpyxl import load_workbook
    table = '../data/ukpopestimatesmid2020on2021geography.xls'
    ws = load_workbook(table)['Tables']
    print({key : value for key, value in ws.tables.items()})

    return pd.read_excel(table,sheet_name=1)






# lets use the nomis output for 2011 then...
''' THIS STILL IS NOT A REAL CSV FILE!'''
population = pd.read_csv('../data/1085757477935590.csv', header=6,index_col=0).dropna()['All usual residents']

@np.vectorize
def populate(cd):
    try: p = population[cd]
    except Exception as e:
        #  print(e)
         p = -1
    finally: return int(p)

for f in files: 
    if not '21' in f:
        continue 

      
    df = pd.read_csv(f,compression='gzip',index_col=0)

    df['population'] = populate(df.index)


    # precission
    df.lng = df.lng.astype('float16')
    df.lat = df.lat.astype('float16')
    df.population.astype('int16')





    df.to_csv(f.split('/')[-1],compression='gzip')

    df.to_csv(f.split('/')[-1].strip('.gz'))

    df.reset_index().to_feather(f.split('/')[-1].replace('.csv.gz','.feather'),compression="lz4")





    print(f)